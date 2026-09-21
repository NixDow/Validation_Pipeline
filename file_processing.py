import os
from datetime import datetime
from openpyxl import load_workbook, Workbook
from openpyxl.styles import PatternFill, Alignment, Font
from config import (VALIDATION_RULES,
                    BLANK_FILL,
                    INVALID_FILL,
                    MISMATCH_FILL,
                    QTY_LIMIT,
                    OUTPUT_DIR,
                    REPORT_OUTPUT_DIR,
)
from validator import (
    is_blank,
    val_and_form_date,
    in_date_range,
    make_number,
    apply_conditions,
    drop_downs_list,
    standardise_value,
    valid_email,
    clean_order_id
)
from detect_and_rename import map_columns, detect_header_row, get_max_row

# Initialize fills
blank_fill = PatternFill(start_color=BLANK_FILL, fill_type="solid")
invalid_fill = PatternFill(start_color=INVALID_FILL, fill_type="solid")
mismatch_fill = PatternFill(start_color=MISMATCH_FILL, fill_type="solid")

#how each conditional error type is labelled: (fill, hex, message)
CONDITIONAL_LABELS = {
    "required_if": (blank_fill, BLANK_FILL, "Missing mandatory information; required by another column"),
    "blanks_or_zero_list": (invalid_fill, INVALID_FILL, "Entry should be blank or 0 for this status"),
    "dropdown_mismatch": (mismatch_fill, MISMATCH_FILL, "Selection does not match the parent column's options"),
    "date_order": (invalid_fill, INVALID_FILL, "Date is earlier than the date it depends on"),
}

def save_report(report_data, checked_log, filename):
    """
    Saves an Excel report with two sheets:
    1. Validation Summary: Detailed errors with color coding.
    2. Checked Columns: A log of every column validated.
    """
    os.makedirs(REPORT_OUTPUT_DIR, exist_ok=True)

    # Initialize workbook and the first sheet
    report_wb = Workbook()
    report_ws = report_wb.active
    report_ws.title = "Validation Summary"

    # --- SHEET 1: Validation Summary ---
    report_ws.append(["Tab Name", "Column Name", "Row", "Issue Description"])

    # Process 5-element tuples: (Tab, Column, Row, Message, Hex_Color)
    for error_entry in sorted(report_data):
        report_ws.append(error_entry[:4])

        # Apply the fill to the description cell of the row just created
        hex_color = error_entry[4]
        desc_cell = report_ws.cell(row=report_ws.max_row, column=4)
        desc_cell.fill = PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")

    # --- SHEET 2: Checked Columns Log ---
    log_ws = report_wb.create_sheet(title="Checked Columns")
    log_ws.append(["Tab Name", "Column Name", "Matched Header", "Checked"])

    for tab_name, col_name, actual_col in checked_log:
        log_ws.append([tab_name, col_name, actual_col, "Yes"])

    # Make headers bold for both sheets
    for ws in [report_ws, log_ws]:
        for cell in ws[1]:
            cell.font = Font(bold=True)

    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    report_path = os.path.join(REPORT_OUTPUT_DIR, f"Validation_Report_{filename}_{timestamp}.xlsx")
    report_wb.save(report_path)

    return report_path

def validate_file(filepath):
    #keep_vba = True only for xlsm files
    is_xlsm = str(filepath).lower().endswith('.xlsm')

    # Load the worksheet (keep_vba=True preserves macros if present)
    wb = load_workbook(filepath, keep_vba=is_xlsm)
    report_data = set()
    column_log = []
    active_rules = VALIDATION_RULES
    global_cond_rules = active_rules.get("conditional_rules", [])

    for sheet_name, rules in active_rules.items():
        if sheet_name == "conditional_rules" or sheet_name not in wb.sheetnames:
            continue

        ws = wb[sheet_name] #the worksheet

        # Store Details check - single cell B3
        if sheet_name == "Store Details":
            target_cell = ws.cell(row=3, column=2)
            target_cell.fill = PatternFill(fill_type=None)

            if is_blank(target_cell.value):
                target_cell.fill = blank_fill
                report_data.add((sheet_name, "Store Name (B3)", 3, "Missing mandatory information, cell is blank", BLANK_FILL))
            else:
                target_cell.value = standardise_value(target_cell.value, "Stores")
                if not drop_downs_list(target_cell.value, "Stores"):
                    target_cell.fill = mismatch_fill
                    report_data.add((sheet_name, "Store Name (B3)", 3, "Selection does not match the dropdown options", MISMATCH_FILL))
            continue

        true_max = get_max_row(ws) # all rows in worksheet with an input, skips empty rows

        expected_cols = list(rules["expected_columns"].keys())
        header_row = detect_header_row(ws, expected_cols)

        if not header_row or true_max == 0:
            print(f"No header row detected for {sheet_name}") #idenitfy header rows
            continue

        headers = [cell.value for cell in ws[header_row]]
        column_map = map_columns(expected_cols, headers) #Map expected column names to actual headers

        missing = [c for c in expected_cols if c not in column_map]
        if missing:
            print(f"{os.path.basename(filepath)} [{sheet_name}] columns not found: {missing}")

        #Order ID cleaning / salvaging pass (runs before validation so salvaged IDs aren't flagged as blank)
        id_col = column_map.get("Order ID")
        notes_col = column_map.get("Notes")
        id_index = headers.index(id_col) + 1 if id_col else None
        notes_index = headers.index(notes_col) + 1 if notes_col else None

        if id_index:
            for row in range(header_row + 1, true_max + 1):
                id_cell = ws.cell(row=row, column=id_index)
                notes_val = ws.cell(row=row, column=notes_index).value if notes_index else None
                cleaned_id = clean_order_id(id_cell.value, notes_val)
                if cleaned_id != id_cell.value:
                    id_cell.value = cleaned_id

        #Main validation loop
        for expected_col, actual_col in column_map.items():

            column_log.append((sheet_name, expected_col, actual_col))
            validations = rules["expected_columns"][expected_col]
            col_index = headers.index(actual_col) + 1

            for row in range(header_row + 1, true_max + 1):
                cell = ws.cell(row=row, column=col_index)
                value = cell.value #extract the cell we need

                def flag(fill, hex_color, message):
                    cell.fill = fill
                    report_data.add((sheet_name, expected_col, row, message, hex_color))

                # Reset any original colour formatting
                cell.fill = PatternFill(fill_type=None)

                #if the cell has an indent in it then remove it where necessary
                if cell.alignment.indent > 0:
                    cell.alignment = Alignment(indent=0)

                #handle number conversions for related rules
                if "number" in validations and not is_blank(value):
                    converted = make_number(value)

                    if converted == 'invalid': #if the make_number functions returns invalid, label the value as invalid
                        flag(invalid_fill, INVALID_FILL, "Invalid number format")
                        continue
                    cell.value = converted
                    value = converted

                #Quantity limit
                if "qty_limit" in validations and isinstance(value, (int, float)):
                    if value > QTY_LIMIT or value < 0:
                        flag(invalid_fill, INVALID_FILL, f"Quantity outside the 0-{QTY_LIMIT} range")
                        continue

                # Mandatory Blank check
                if "not_blank" in validations and is_blank(value):
                    flag(blank_fill, BLANK_FILL, "Missing mandatory information; cell is blank.")
                    continue

                # Date Validation & Formatting
                if "date" in validations:
                    status, parsed_date = val_and_form_date(value)

                    if status == 'valid':
                        # Standardize to YYYY/MM/DD
                        cell.value = parsed_date
                        if "date_range" in validations and not in_date_range(parsed_date):
                            flag(invalid_fill, INVALID_FILL, "Date is outside the plausible range")
                            continue
                    elif status == 'invalid':
                        flag(invalid_fill, INVALID_FILL, "Invalid date format submitted.")
                        continue

                #Dropdown standardisation + validation
                for rule in validations:
                    if rule.startswith("list"):
                        list_name = rule.split(":")[1]  #extract the dropdown name

                        standardised = standardise_value(value, list_name)
                        if standardised != value:
                            value = standardised
                            cell.value = value

                        if not drop_downs_list(value, list_name): #if the input isn't in the list
                            flag(mismatch_fill, MISMATCH_FILL, "Selection does not match the dropdown options.")

                if "email" in validations and not valid_email(value):
                    flag(invalid_fill, INVALID_FILL, "Incorrect email format")

        #Conditional validation - once per row, after every column has been cleaned
        if global_cond_rules:
            for row in range(header_row + 1, true_max + 1):
                conditional_errors = apply_conditions(ws, row, column_map, global_cond_rules, headers)
                for err_col, error_type in conditional_errors:
                    # The column_map already handles whether the column exists in this sheet
                    if err_col in column_map:
                        err_col_idx = headers.index(column_map[err_col]) + 1
                        fill, hex_color, message = CONDITIONAL_LABELS[error_type]
                        ws.cell(row=row, column=err_col_idx).fill = fill
                        report_data.add((sheet_name, err_col, row, message, hex_color))

    filename = os.path.basename(filepath)
    if report_data or column_log:
        save_report(report_data, column_log, filename)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    wb.save(os.path.join(OUTPUT_DIR, filename))

    return len(report_data)
