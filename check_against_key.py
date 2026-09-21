
import csv
from openpyxl import load_workbook
from config import INPUT_FOLDER, OUTPUT_DIR, VALIDATION_RULES
from detect_and_rename import detect_header_row, map_columns

def cell_fill(cell):
    rgb = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type else None
    return rgb[-6:] if isinstance(rgb, str) else None

def main():
    with open(INPUT_FOLDER / "answer_key.csv", encoding="utf-8") as f:
        key = list(csv.DictReader(f))

    books, maps = {}, {}
    passed, failures = 0, []
    expected_cells = set()

    for entry in key:
        filename, tab, row = entry["File"], entry["Tab"], int(entry["Row"])
        if filename not in books:
            books[filename] = load_workbook(OUTPUT_DIR / filename)
            ws = books[filename]["Orders"]
            expected_cols = list(VALIDATION_RULES["Orders"]["expected_columns"])
            header_row = detect_header_row(ws, expected_cols)
            headers = [c.value for c in ws[header_row]]
            maps[filename] = (header_row, headers, map_columns(expected_cols, headers))

        ws = books[filename][tab]
        if tab == "Store Details":
            cell = ws["B3"]
        else:
            _, headers, column_map = maps[filename]
            cell = ws.cell(row=row, column=headers.index(column_map[entry["Column"]]) + 1)
        expected_cells.add((filename, tab, cell.coordinate))

        outcome = entry["Expected Outcome"]
        fill = cell_fill(cell)
        if outcome.startswith("flagged"):
            ok = fill == outcome.split()[1]
        else:
            want = outcome.split("-> ", 1)[1]
            got = cell.value
            ok = fill is None and (str(got) == want or (isinstance(got, (int, float)) and got == float(want)))

        if ok:
            passed += 1
        else:
            failures.append(f"{filename} {tab}!{cell.coordinate} [{entry['Fault']}] "
                            f"injected={entry['Injected Value']!r} expected '{outcome}' got value={cell.value!r} fill={fill}")

    #false positives: filled cells in data rows that weren't injected
    false_positives = []
    for filename, wb in books.items():
        header_row, _, _ = maps[filename]
        for row in wb["Orders"].iter_rows(min_row=header_row + 1):
            for cell in row:
                if cell_fill(cell) and (filename, "Orders", cell.coordinate) not in expected_cells:
                    false_positives.append(f"{filename} Orders!{cell.coordinate} value={cell.value!r} fill={cell_fill(cell)}")

    print(f"{passed}/{len(key)} injected faults handled as expected")
    for f in failures:
        print("  FAIL", f)
    print(f"{len(false_positives)} unexpected flags")
    for fp in false_positives:
        print("  EXTRA", fp)

if __name__ == "__main__":
    main()
